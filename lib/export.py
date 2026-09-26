"""Render the structured resume out to a real file: .md, .txt, .docx, .tex, .pdf."""

import re
import shutil
import subprocess
from pathlib import Path

from lib import templates as _templates
from lib import latex as _latex
from lib.errors import ResumeTailorError
from lib.ids import cert_text, skill_names
from lib.links import linkedin_label, normalize_url
from lib.locking import sha256_bytes

TECTONIC_BIN = Path(__file__).resolve().parent.parent / "bin" / "tectonic"

# Original, template-agnostic ordering -- used when no template_id is given
# so existing callers keep their exact prior behavior.
_DEFAULT_SECTION_ORDER = ["summary", "skills", "experience", "education", "projects", "certifications"]


def _section_order(template_id: str | None) -> list[str]:
    if not template_id:
        return _DEFAULT_SECTION_ORDER
    tmpl = _templates.get_template(template_id)
    if tmpl is None:
        raise ValueError(f"Unknown template '{template_id}'. Use list_templates to see valid ids.")
    # Layouts may reference sections (e.g. "achievements") that aren't part
    # of the resume schema -- drop anything we don't know how to render.
    return [s for s in tmpl["layout"]["sections"] if s in _DEFAULT_SECTION_ORDER]


def _bullets_text(bullets) -> list[str]:
    return [b.get("text", "") if isinstance(b, dict) else str(b) for b in bullets or []]


def _md_summary(resume: dict) -> list[str]:
    return ["## Summary", resume["summary"], ""] if resume.get("summary") else []


def _md_skills(resume: dict) -> list[str]:
    if not resume.get("skills"):
        return []
    lines = ["## Skills"]
    for group in resume["skills"]:
        items = ", ".join(skill_names(group))
        lines.append(f"- {group.get('category', 'General')}: {items}")
    lines.append("")
    return lines


def _md_experience(resume: dict) -> list[str]:
    if not resume.get("experience"):
        return []
    lines = ["## Experience"]
    for exp in resume["experience"]:
        header = f"### {exp.get('title', '')}, {exp.get('company', '')} | {exp.get('location', '')} | {exp.get('start', '')} - {exp.get('end', '')}"
        lines.append(header)
        for b in _bullets_text(exp.get("bullets")):
            lines.append(f"- {b}")
        lines.append("")
    return lines


def _md_education(resume: dict) -> list[str]:
    if not resume.get("education"):
        return []
    lines = ["## Education"]
    for edu in resume["education"]:
        lines.append(f"### {edu.get('degree', '')}, {edu.get('school', '')} | {edu.get('year', '')}")
    lines.append("")
    return lines


def _md_projects(resume: dict) -> list[str]:
    if not resume.get("projects"):
        return []
    lines = ["## Projects"]
    for proj in resume["projects"]:
        name = proj.get("name", "")
        github = proj.get("github")
        heading = f"[{name}]({normalize_url(github)})" if github else name
        lines.append(f"### {heading}")
        for b in _bullets_text(proj.get("bullets")):
            lines.append(f"- {b}")
    lines.append("")
    return lines


def _md_certifications(resume: dict) -> list[str]:
    if not resume.get("certifications"):
        return []
    lines = ["## Certifications"]
    for cert in resume["certifications"]:
        lines.append(f"- {cert_text(cert)}")
    lines.append("")
    return lines


_MD_SECTION_RENDERERS = {
    "summary": _md_summary,
    "skills": _md_skills,
    "experience": _md_experience,
    "education": _md_education,
    "projects": _md_projects,
    "certifications": _md_certifications,
}


def _md_contact_line(contact: dict) -> str:
    parts = []
    if contact.get("phone"):
        parts.append(contact["phone"])
    if contact.get("email"):
        parts.append(f"[{contact['email']}](mailto:{contact['email']})")
    if contact.get("website"):
        url = contact["website"]
        parts.append(f"[{url}]({normalize_url(url)})")
    if contact.get("linkedin"):
        url = contact["linkedin"]
        parts.append(f"[{linkedin_label(url)}]({normalize_url(url)})")
    if contact.get("github"):
        url = contact["github"]
        parts.append(f"[{url}]({normalize_url(url)})")
    if contact.get("substack"):
        url = contact["substack"]
        parts.append(f"[{url}]({normalize_url(url)})")
    if contact.get("location"):
        parts.append(contact["location"])
    return " | ".join(parts)


def to_markdown(resume: dict, template_id: str | None = None) -> str:
    lines = [f"# {resume.get('name', '')}"]
    contact_line = _md_contact_line(resume.get("contact", {}) or {})
    if contact_line:
        lines.append(contact_line)
    lines.append("")

    for section in _section_order(template_id):
        lines += _MD_SECTION_RENDERERS[section](resume)

    return "\n".join(lines).strip() + "\n"


_MD_LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


def _flatten_md_link(match: re.Match) -> str:
    text, url = match.group(1), match.group(2)
    bare = url.removeprefix("mailto:").removeprefix("https://").removeprefix("http://")
    # Skip the redundant "(url)" when the visible text already *is* the URL
    # (plain email/website/github links) -- keep it for short labels like
    # the "in/username" LinkedIn label or a project name, where it adds info.
    return text if text == bare else f"{text} ({url})"


def to_txt(resume: dict, template_id: str | None = None) -> str:
    # Plain text is the safest possible ATS format -- strip markdown symbols.
    md = to_markdown(resume, template_id)
    lines = []
    for line in md.splitlines():
        line = _MD_LINK_RE.sub(_flatten_md_link, line)
        lines.append(line.lstrip("#").lstrip("-").strip() if line.strip() else "")
    return "\n".join(lines)


# Templates whose source layout uses colored section headers (see
# data/templates/templates.yaml) get that color in the .docx export too.
_HEADING_COLOR_BY_TEMPLATE = {"full-stack-modern": (0x1F, 0x4E, 0x79)}


def _add_hyperlink(paragraph, url: str, text: str, bold: bool = False):
    """Insert a real clickable hyperlink run into a docx paragraph.
    python-docx has no built-in hyperlink API, so this builds the
    w:hyperlink OOXML element directly."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE

    part = paragraph.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    if bold:
        rpr.append(OxmlElement("w:b"))
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    run.append(rpr)

    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def _docx_contact_line(document, contact: dict) -> None:
    parts: list[tuple[str, str, str | None]] = []
    if contact.get("phone"):
        parts.append(("text", contact["phone"], None))
    if contact.get("email"):
        parts.append(("link", contact["email"], f"mailto:{contact['email']}"))
    if contact.get("website"):
        url = contact["website"]
        parts.append(("link", url, normalize_url(url)))
    if contact.get("linkedin"):
        url = contact["linkedin"]
        parts.append(("link", linkedin_label(url), normalize_url(url)))
    if contact.get("github"):
        url = contact["github"]
        parts.append(("link", url, normalize_url(url)))
    if contact.get("substack"):
        url = contact["substack"]
        parts.append(("link", url, normalize_url(url)))
    if contact.get("location"):
        parts.append(("text", contact["location"], None))

    if not parts:
        return
    p = document.add_paragraph()
    for i, (kind, label, href) in enumerate(parts):
        if i > 0:
            p.add_run("  |  ")
        if kind == "link":
            _add_hyperlink(p, href, label)
        else:
            p.add_run(label)


def _docx_heading(document, text: str, template_id: str | None):
    heading = document.add_heading(text, level=2)
    color = _HEADING_COLOR_BY_TEMPLATE.get(template_id)
    if color:
        from docx.shared import RGBColor
        for run in heading.runs:
            run.font.color.rgb = RGBColor(*color)


def _docx_summary(document, resume: dict, template_id: str | None) -> None:
    if resume.get("summary"):
        _docx_heading(document, "Summary", template_id)
        document.add_paragraph(resume["summary"])


def _docx_skills(document, resume: dict, template_id: str | None) -> None:
    if resume.get("skills"):
        _docx_heading(document, "Skills", template_id)
        for group in resume["skills"]:
            items = ", ".join(skill_names(group))
            document.add_paragraph(f"{group.get('category', 'General')}: {items}", style="List Bullet")


def _docx_experience(document, resume: dict, template_id: str | None) -> None:
    if resume.get("experience"):
        _docx_heading(document, "Experience", template_id)
        for exp in resume["experience"]:
            p = document.add_paragraph()
            p.add_run(f"{exp.get('title', '')}, {exp.get('company', '')}").bold = True
            meta = " | ".join(v for v in [exp.get("location"), f"{exp.get('start', '')} - {exp.get('end', '')}"] if v)
            if meta:
                document.add_paragraph(meta)
            for b in _bullets_text(exp.get("bullets")):
                document.add_paragraph(b, style="List Bullet")


def _docx_education(document, resume: dict, template_id: str | None) -> None:
    if resume.get("education"):
        _docx_heading(document, "Education", template_id)
        for edu in resume["education"]:
            document.add_paragraph(f"{edu.get('degree', '')}, {edu.get('school', '')} ({edu.get('year', '')})")


def _docx_projects(document, resume: dict, template_id: str | None) -> None:
    if resume.get("projects"):
        _docx_heading(document, "Projects", template_id)
        for proj in resume["projects"]:
            p = document.add_paragraph()
            github = proj.get("github")
            if github:
                _add_hyperlink(p, normalize_url(github), proj.get("name", ""), bold=True)
            else:
                p.add_run(proj.get("name", "")).bold = True
            for b in _bullets_text(proj.get("bullets")):
                document.add_paragraph(b, style="List Bullet")


def _docx_certifications(document, resume: dict, template_id: str | None) -> None:
    if resume.get("certifications"):
        _docx_heading(document, "Certifications", template_id)
        for cert in resume["certifications"]:
            document.add_paragraph(cert_text(cert), style="List Bullet")


_DOCX_SECTION_RENDERERS = {
    "summary": _docx_summary,
    "skills": _docx_skills,
    "experience": _docx_experience,
    "education": _docx_education,
    "projects": _docx_projects,
    "certifications": _docx_certifications,
}


def to_docx(resume: dict, out_path: str, template_id: str | None = None) -> str:
    try:
        import docx
        from docx.shared import Pt
    except ImportError as e:
        raise RuntimeError("Exporting to .docx requires python-docx. Install it with: pip install python-docx") from e

    order = _section_order(template_id)  # raises if template_id is unknown

    document = docx.Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    document.add_heading(resume.get("name", ""), level=1)
    _docx_contact_line(document, resume.get("contact", {}) or {})

    for section in order:
        _DOCX_SECTION_RENDERERS[section](document, resume, template_id)

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    document.save(out_path)
    return out_path


def to_tex(resume: dict, template_id: str | None = None) -> str:
    """Render the resume into the Jake's-Resume-style LaTeX source (see
    lib/latex.py). This is the same source the PDF is compiled from, so the
    two outputs can never drift apart."""
    return _latex.render_latex(resume, template_id)


def _tectonic_path() -> str:
    if TECTONIC_BIN.exists():
        return str(TECTONIC_BIN)
    found = shutil.which("tectonic")
    if found:
        return found
    raise ResumeTailorError(
        "LATEX_COMPILE_FAILED",
        "No tectonic binary found (expected at bin/tectonic, or 'tectonic' on "
        "PATH). Download the standalone release for your platform from "
        "https://github.com/tectonic-typesetting/tectonic/releases and place "
        "it at bin/tectonic.",
    )


_LOG_MAX_CHARS = 20_000
_LOG_TAIL_CHARS = 2_000
_BASENAME_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._-]{0,127}$")
_OVERFULL_RE = re.compile(r"Overfull \\hbox")
_UNDERFULL_RE = re.compile(r"Underfull \\hbox")


def _validate_basename(basename: str) -> str:
    if not isinstance(basename, str) or "/" in basename or "\\" in basename or ".." in basename:
        raise ResumeTailorError("PATH_TRAVERSAL", "basename must be a plain filename stem (no path separators or '..').")
    if not _BASENAME_RE.match(basename):
        raise ResumeTailorError("INVALID_ID", "basename may only contain letters, digits, '.', '_' and '-'.")
    return basename


def _last_tex_pass(log: str) -> str:
    """Tectonic reruns TeX until the aux files settle and repeats every box
    warning on each pass -- count only the final pass."""
    cut = max(log.rfind("note: Rerunning TeX"), log.rfind("note: Running TeX"))
    return log[cut:] if cut >= 0 else log


def compile_tex(tex: str, out_dir: Path, basename: str, timeout: int = 120) -> dict:
    """Compile exactly `tex` (spec §44). The string is written byte-for-byte
    to `out_dir/<basename>.tex`, compiled with tectonic, and the written file
    is re-hashed afterwards so the returned `tex_sha256` provably identifies
    the source the PDF was built from."""
    _validate_basename(basename)
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    tex_path = out_dir / f"{basename}.tex"
    pdf_path = out_dir / f"{basename}.pdf"

    tex_bytes = tex.encode("utf-8")
    expected_sha = sha256_bytes(tex_bytes)
    tex_path.write_bytes(tex_bytes)
    if pdf_path.exists():
        pdf_path.unlink()  # never report a stale PDF from an earlier run

    tectonic = _tectonic_path()
    try:
        result = subprocess.run(
            [tectonic, "--outdir", str(out_dir), str(tex_path)],
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,  # one interleaved stream, so pass boundaries stay in order
            text=True,
            timeout=timeout,
        )
    except subprocess.TimeoutExpired as e:
        raise ResumeTailorError(
            "LATEX_COMPILE_FAILED",
            f"tectonic timed out after {timeout}s.",
            details={"log_tail": ""},
        ) from e

    full_log = result.stdout or ""
    if result.returncode != 0 or not pdf_path.exists():
        raise ResumeTailorError(
            "LATEX_COMPILE_FAILED",
            f"tectonic compilation failed (exit code {result.returncode}).",
            details={"log_tail": full_log[-_LOG_TAIL_CHARS:]},
        )

    actual_sha = sha256_bytes(tex_path.read_bytes())
    if actual_sha != expected_sha:
        raise ResumeTailorError(
            "LATEX_COMPILE_FAILED",
            "The .tex file on disk no longer matches the compiled source.",
        )

    last_pass = _last_tex_pass(full_log)
    return {
        "pdf_path": str(pdf_path),
        "tex_path": str(tex_path),
        "log": full_log[-_LOG_MAX_CHARS:],
        "tex_sha256": expected_sha,
        "pdf_sha256": sha256_bytes(pdf_path.read_bytes()),
        "overfull_hbox_count": len(_OVERFULL_RE.findall(last_pass)),
        "underfull_hbox_count": len(_UNDERFULL_RE.findall(last_pass)),
    }


def to_pdf(resume: dict, out_path: str, template_id: str | None = None) -> str:
    """Render the resume to LaTeX once and compile that exact string with
    tectonic (a self-contained LaTeX engine bundled at bin/tectonic). The
    .tex is left next to the PDF (same stem)."""
    out_path = Path(out_path)
    info = compile_tex(to_tex(resume, template_id), out_path.parent, out_path.stem)
    produced = Path(info["pdf_path"])
    if produced != out_path:
        produced.replace(out_path)
    return str(out_path)


def export_resume(resume: dict, out_path: str, fmt: str, template_id: str | None = None) -> str:
    fmt = fmt.lower().lstrip(".")
    out_path = str(out_path)
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)

    if fmt == "md":
        Path(out_path).write_text(to_markdown(resume, template_id), encoding="utf-8")
        return out_path
    if fmt == "txt":
        Path(out_path).write_text(to_txt(resume, template_id), encoding="utf-8")
        return out_path
    if fmt == "docx":
        return to_docx(resume, out_path, template_id)
    if fmt == "tex":
        Path(out_path).write_text(to_tex(resume, template_id), encoding="utf-8")
        return out_path
    if fmt == "pdf":
        return to_pdf(resume, out_path, template_id)
    raise ValueError(f"Unsupported export format: {fmt} (use md, txt, docx, tex, or pdf)")
